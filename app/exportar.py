import csv
import io
from datetime import date

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIPOS = {"xlsx": XLSX, "csv": "text/csv; charset=utf-8", "pdf": "application/pdf"}


def ficheiro(colunas: list[tuple[str, str]], linhas: list[dict], nome: str, formato: str = "xlsx") -> tuple[str, bytes, str]:
    """colunas: [(chave, cabeçalho)]. Devolve (nome do ficheiro, conteúdo, tipo MIME) em Excel ou CSV."""
    nome_ficheiro = f"{nome}-{date.today():%Y%m%d}.{formato}"
    if formato == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow([c[1] for c in colunas])
        for l in linhas:
            w.writerow([l.get(c[0]) for c in colunas])
        return nome_ficheiro, ("﻿" + buf.getvalue()).encode("utf-8"), TIPOS["csv"]

    wb = Workbook()
    ws = wb.active
    ws.title = nome[:31]
    ws.append([c[1] for c in colunas])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for l in linhas:
        ws.append([l.get(c[0]) for c in colunas])
    for i, (chave, cab) in enumerate(colunas, start=1):
        largura = max([len(str(cab))] + [len(str(l.get(chave) or "")) for l in linhas[:200]])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(45, largura + 2)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return nome_ficheiro, out.getvalue(), XLSX


def tabela(colunas: list[tuple[str, str]], linhas: list[dict], nome: str, formato: str = "xlsx") -> Response:
    """O mesmo ficheiro, como resposta HTTP para descarregar."""
    nome_ficheiro, dados, mime = ficheiro(colunas, linhas, nome, formato)
    return Response(dados, media_type=mime, headers={"Content-Disposition": f'attachment; filename="{nome_ficheiro}"'})
